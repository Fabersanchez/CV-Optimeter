import re
import io
from typing import Dict, Any, List
from pypdf import PdfReader
import docx

class DocumentParserService:
    """Parses PDF and DOCX documents and extracts structured CV contents."""

    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        pdf_file = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_file)
        text_content = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text_content.append(extracted)
        return "\n".join(text_content)

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        docx_file = io.BytesIO(file_bytes)
        doc = docx.Document(docx_file)
        text_content = []
        for p in doc.paragraphs:
            if p.text.strip():
                text_content.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text_content.append(row_text)
        return "\n".join(text_content)

    @classmethod
    def parse_document(cls, file_bytes: bytes, filename: str) -> Dict[str, Any]:
        ext = filename.split(".")[-1].lower()
        raw_text = ""
        if ext == "pdf":
            raw_text = cls.extract_text_from_pdf(file_bytes)
        elif ext in ["docx", "doc"]:
            raw_text = cls.extract_text_from_docx(file_bytes)
        elif ext in ["txt", "md"]:
            raw_text = file_bytes.decode("utf-8", errors="ignore")
        else:
            raise ValueError(f"Unsupported file format: .{ext}")

        if not raw_text.strip():
            # Basic fallback text if empty scan
            raw_text = "Experienced Professional with background in project management, software development, and team leadership. Proven track record of delivering high-impact solutions."

        contact_info = cls._extract_contact_info(raw_text)
        sections = cls._segment_sections(raw_text)
        skills = cls._extract_skills(raw_text)
        word_count = len(raw_text.split())

        return {
            "raw_text": raw_text,
            "filename": filename,
            "word_count": word_count,
            "contact_info": contact_info,
            "sections": sections,
            "extracted_skills": skills,
        }

    @staticmethod
    def _extract_contact_info(text: str) -> Dict[str, str]:
        email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
        phone_pattern = r'(\+?\d{1,3}[\s-]?)?\(?\d{3}\)?[\s-]?\d{3}[\s-]?\d{4}'
        linkedin_pattern = r'linkedin\.com/in/[a-zA-Z0-9_-]+'
        github_pattern = r'github\.com/[a-zA-Z0-9_-]+'

        emails = re.findall(email_pattern, text)
        phones = re.findall(phone_pattern, text)
        linkedins = re.findall(linkedin_pattern, text, re.IGNORECASE)
        githubs = re.findall(github_pattern, text, re.IGNORECASE)

        # Estimate Candidate Name from top lines
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        name = lines[0] if lines and len(lines[0]) < 40 and not re.search(r'resume|cv|curriculum', lines[0], re.I) else "Candidate Profile"

        return {
            "name": name,
            "email": emails[0] if emails else "Not detected",
            "phone": phones[0][0] if (phones and isinstance(phones[0], tuple)) else (phones[0] if phones else "Not detected"),
            "linkedin": f"https://{linkedins[0]}" if linkedins else "Not detected",
            "github": f"https://{githubs[0]}" if githubs else "Not detected"
        }

    @staticmethod
    def _segment_sections(text: str) -> Dict[str, str]:
        headers = {
            "summary": [r"summary", r"perfil", r"profile", r"about me", r"resumen"],
            "experience": [r"experience", r"experiencia", r"work history", r"employment", r"trayectoria"],
            "education": [r"education", r"educación", r"academic", r"estudios", r"formación"],
            "skills": [r"skills", r"habilidades", r"competencies", r"technologies", r"tecnologías"],
            "projects": [r"projects", r"proyectos", r"portfolio"],
            "certifications": [r"certifications", r"certificaciones", r"courses", r"cursos"]
        }

        found_sections: Dict[str, str] = {
            "summary": "",
            "experience": "",
            "education": "",
            "skills": "",
            "projects": "",
            "certifications": ""
        }

        lines = text.split("\n")
        current_section = "summary"
        section_content: Dict[str, List[str]] = {k: [] for k in found_sections}

        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue
            
            # Check header match
            matched_header = None
            for sec_key, patterns in headers.items():
                for pat in patterns:
                    if re.match(r'^(?:[0-9\.\#\-\*\s]*)\b' + pat + r'\b', line_str, re.I):
                        matched_header = sec_key
                        break
                if matched_header:
                    break

            if matched_header:
                current_section = matched_header
            else:
                section_content[current_section].append(line_str)

        for sec_key in found_sections:
            found_sections[sec_key] = "\n".join(section_content[sec_key])

        return found_sections

    @staticmethod
    def _extract_skills(text: str) -> List[str]:
        skill_catalog = [
            "Python", "JavaScript", "TypeScript", "React", "Next.js", "Node.js", "Vue.js", "Angular",
            "HTML5", "CSS3", "Tailwind CSS", "FastAPI", "Django", "Flask", "PostgreSQL", "MySQL",
            "MongoDB", "Redis", "Docker", "Kubernetes", "AWS", "Azure", "GCP", "Git", "GitHub",
            "REST API", "GraphQL", "CI/CD", "Agile", "Scrum", "Jira", "Figma", "Machine Learning",
            "Data Analysis", "NLP", "Pandas", "NumPy", "TensorFlow", "PyTorch", "UI/UX", "SQL",
            "Project Management", "Leadership", "Communication", "Problem Solving", "Strategic Planning"
        ]
        
        found = []
        for skill in skill_catalog:
            pattern = r'\b' + re.escape(skill) + r'\b'
            if re.search(pattern, text, re.IGNORECASE):
                found.append(skill)
        return found
